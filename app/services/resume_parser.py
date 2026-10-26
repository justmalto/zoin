"""Bounded extraction for the supported resume formats."""
import io
import re
from zipfile import ZipFile, BadZipFile
from docx import Document
from pypdf import PdfReader

MAX_TEXT_CHARACTERS = 18000
MAX_PDF_PAGES = 50
MAX_DOCX_EXPANDED_BYTES = 20 * 1024 * 1024


def bounded_text(parts):
    result=[]; length=0
    for part in parts:
        length += len(part) + 1
        if length > MAX_TEXT_CHARACTERS + 1:
            raise ValueError('Resume text exceeds 18,000 characters. Upload a shorter resume.')
        result.append(part)
    return '\n'.join(result).strip()


def extract_text(file_bytes: bytes, content_type: str, file_name: str) -> str:
    # Dispatch by the validated MIME type, never a conflicting filename suffix.
    if content_type == 'application/pdf':
        if not file_bytes.startswith(b'%PDF-'):
            raise ValueError('Upload a valid PDF resume.')
        reader=PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            raise ValueError('Upload an unencrypted PDF resume.')
        if len(reader.pages)>MAX_PDF_PAGES:
            raise ValueError('PDF resumes must contain at most 50 pages.')
        return bounded_text(page.extract_text() or '' for page in reader.pages)
    if content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        try:
            with ZipFile(io.BytesIO(file_bytes)) as archive:
                entries=archive.infolist()
                if len(entries)>1000 or sum(item.file_size for item in entries)>MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError('DOCX resume expands beyond the supported document limit.')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError('Upload a valid DOCX resume.')
        except BadZipFile:
            raise ValueError('Upload a valid DOCX resume.') from None
        document=Document(io.BytesIO(file_bytes))
        return bounded_text([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    if len(file_bytes)>MAX_TEXT_CHARACTERS*4:
        raise ValueError('Resume text exceeds 18,000 characters. Upload a shorter resume.')
    text=file_bytes.decode('utf-8-sig').strip()
    if '\x00' in text:
        raise ValueError('Upload a UTF-8 text resume.')
    return bounded_text([text])


def chunk_text(text: str, max_characters: int = 1100) -> list[dict]:
    heading_pattern=re.compile(r'(?im)^(summary|profile|experience|employment|education|skills|projects|certifications|awards)\s*$')
    section='general';chunks=[]
    for part in heading_pattern.split(text):
        cleaned=part.strip()
        if not cleaned:continue
        if heading_pattern.fullmatch(cleaned):section=cleaned.lower();continue
        for start in range(0,len(cleaned),max_characters):
            content=cleaned[start:start+max_characters].strip()
            if content:chunks.append({'section':section,'chunk_index':len(chunks),'content':content})
    return chunks or [{'section':'general','chunk_index':0,'content':text}]
