import io,sys,unittest
from pathlib import Path
from unittest.mock import Mock,patch
from zipfile import ZipFile,ZIP_DEFLATED
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
from fastapi import Request
from fastapi.testclient import TestClient
from docx import Document
from postgrest.exceptions import APIError
from supabase_auth.errors import AuthApiError
from core.auth import require_user
from main import create_app
from services.resume_parser import extract_text
from services.mail_sync import sync_account

class ParserRegressionTests(unittest.TestCase):
 def test_plain_text_with_pdf_filename_is_read_as_plain_text(self):
  self.assertEqual(extract_text(b'Python ML resume','text/plain','resume.pdf'),'Python ML resume')
 def test_tables_in_docx_are_not_silently_lost(self):
  doc=Document();doc.add_table(rows=1,cols=1).cell(0,0).text='Python machine learning experience'
  output=io.BytesIO();doc.save(output)
  self.assertIn('Python machine learning experience',extract_text(output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document','resume.docx'))
 def test_oversized_expanded_docx_is_rejected_before_parser(self):
  output=io.BytesIO()
  with ZipFile(output,'w',ZIP_DEFLATED) as archive:archive.writestr('word/document.xml',b'x'*(21*1024*1024))
  with patch('services.resume_parser.Document') as parse:
   with self.assertRaisesRegex(ValueError,'expands'):extract_text(output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document','resume.docx')
   parse.assert_not_called()
 def test_too_long_and_binary_text_is_rejected(self):
  for content in [b'a'*18001,b'invalid\xfftext',b'Python\x00text']:
   with self.assertRaises(ValueError):extract_text(content,'text/plain','resume.txt')

class AuthAndGenerationTests(unittest.TestCase):
 def test_auth_outage_is_not_reported_as_expired_session(self):
  with patch('core.auth.auth_client') as auth:
   auth.return_value.auth.get_user.side_effect=OSError('private network diagnostic')
   response=TestClient(create_app()).get('/api/v1/profile',headers={'Authorization':'Bearer fixture'})
  self.assertEqual(response.status_code,503);self.assertNotIn('private network',response.text)
 def test_invalid_token_still_returns_401(self):
  with patch('core.auth.auth_client') as auth:
   auth.return_value.auth.get_user.side_effect=AuthApiError('invalid',401,None)
   response=TestClient(create_app()).get('/api/v1/profile',headers={'Authorization':'Bearer fixture'})
  self.assertEqual(response.status_code,401)
 def test_match_uses_atomic_input_guard_and_reports_stale_conflict(self):
  app=create_app()
  def user(request:Request):request.state.access_token='fixture';return 'owner'
  app.dependency_overrides[require_user]=user
  db=Mock()
  def table(name):
   data=[{'id':'job','job_description':'Original sufficiently long job text','selected_resume_id':'resume'}] if name=='applications' else [{'id':'resume','resume_text':'Original sufficiently long resume','content_hash':'original-hash'}]
   query=Mock();query.select.return_value=query;query.eq.return_value=query;query.limit.return_value=query;query.execute.return_value.data=data;return query
  db.table.side_effect=table
  db.rpc.return_value.execute.side_effect=APIError({'message':'inputs changed','code':'40001','details':None,'hint':None})
  analysis={'resume_id':'resume','match_score':75,'model_name':'fixture'}
  with patch('routers.recommendations.user_client',return_value=db),patch('routers.recommendations.generate_recommendations',return_value=[analysis]):
   response=TestClient(app).post('/api/v1/applications/11111111-1111-4111-8111-111111111111/recommendations')
  self.assertEqual(response.status_code,409)
  name,args=db.rpc.call_args.args;self.assertEqual(name,'store_resume_recommendation');self.assertEqual(args['p_expected_hash'],'original-hash')

class SyncBoundaryTests(unittest.TestCase):
 def test_twenty_jobs_can_complete_and_twenty_one_are_bounded(self):
  for count,status in [(20,'success'),(21,'partial')]:
   calls=[]
   def worker(action,**payload):
    calls.append((action,payload))
    if action=='claim':return {'state':'claimed','run_id':'run','email_address':'x@example.invalid','ciphertext':'x'}
    if action=='jobs':return [{'id':str(i)} for i in range(count)]
    return {}
   with patch('services.mail_sync.worker',side_effect=worker),patch('services.mail_sync.decrypt_password',return_value='secret'),patch('services.mail_sync.run_mail_workflow',return_value={'matches':[],'has_more':False}):
    self.assertEqual(sync_account('owner')['status'],status)
   self.assertEqual(sum(a=='record' for a,_ in calls),20)
 def test_database_failure_during_finish_never_escapes_background_task(self):
  def worker(action,**payload):
   if action=='claim':return {'state':'claimed','run_id':'run','email_address':'x','ciphertext':'bad'}
   raise OSError('database unavailable')
  with patch('services.mail_sync.worker',side_effect=worker),patch('services.mail_sync.decrypt_password',side_effect=ValueError('bad envelope')):
   self.assertEqual(sync_account('owner')['status'],'failed')
